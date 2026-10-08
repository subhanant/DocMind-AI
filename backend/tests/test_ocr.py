"""
Tests for DocumentOCRService (Phase 2).
"""
import io
import pytest
import pymupdf

from app.services.ocr_service import DocumentOCRService


def test_plain_text_fallback():
    sample_text = "Sample Bill of Lading text content."
    res = DocumentOCRService.extract_from_bytes(
        file_bytes=sample_text.encode("utf-8"),
        filename="test.txt",
        mime_type="text/plain",
    )
    assert res.full_text == sample_text
    assert res.page_count == 1
    assert res.extractor_used == "plain_text"


def test_pymupdf_pdf_extraction():
    # Create an in-memory PDF using pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "BILL OF LADING\nB/L NO: MSCU8829104\nVESSEL: MSC OSCAR\nPORT OF LOADING: SHANGHAI")
    pdf_bytes = doc.tobytes()
    doc.close()

    res = DocumentOCRService.extract_from_bytes(
        file_bytes=pdf_bytes,
        filename="test_bol.pdf",
        mime_type="application/pdf",
    )
    assert res.page_count == 1
    assert "BILL OF LADING" in res.full_text
    assert "MSCU8829104" in res.full_text
    assert "MSC OSCAR" in res.full_text
    assert res.extractor_used == "pymupdf"
    assert not res.is_scanned


def test_docx_extraction():
    import docx
    doc = docx.Document()
    doc.add_paragraph("BILL OF LADING NO: MAEU9920192")
    doc.add_paragraph("Shipper: ACME Global Trading Corp")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = DocumentOCRService.extract_from_bytes(
        file_bytes=docx_bytes,
        filename="shipping.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "MAEU9920192" in res.full_text
    assert "ACME Global Trading Corp" in res.full_text
    assert res.extractor_used == "python-docx"


def test_xlsx_extraction():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manifest"
    ws.append(["Container No", "Seal No", "Gross Weight"])
    ws.append(["TGHU7654321", "SL98765", "22500 KGS"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    res = DocumentOCRService.extract_from_bytes(
        file_bytes=xlsx_bytes,
        filename="manifest.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert "TGHU7654321" in res.full_text
    assert "SL98765" in res.full_text
    assert res.extractor_used == "openpyxl"
