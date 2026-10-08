"""
Document Text Extraction Pipeline.

Provides clean abstraction DocumentTextExtractor:
- Extracts selectable text from native PDFs via PyMuPDF (fitz)
- Detects scanned / image-only pages and triggers OCR fallback via OCRService
- Extracts text from images (PNG, JPG, JPEG, WEBP) via OCRService
- Preserves explicit page boundaries:
  {
    "pages": [
      {"page_number": 1, "text": "...", "ocr_used": false},
      {"page_number": 2, "text": "...", "ocr_used": true}
    ]
  }
"""
import io
from dataclasses import asdict, dataclass, field
from typing import Any, List, Optional

from app.core.config import get_settings
from app.core.logging import get_logger
from app.services.ocr_service import OCRService, get_ocr_service

logger = get_logger(__name__)
settings = get_settings()


@dataclass
class PageData:
    page_number: int
    text: str
    ocr_used: bool = False
    char_count: int = 0
    lines: List[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "ocr_used": self.ocr_used,
            "char_count": self.char_count,
            "lines": self.lines,
        }


@dataclass
class DocumentPagesResult:
    pages: List[PageData] = field(default_factory=list)
    total_pages: int = 0
    total_chars: int = 0
    is_scanned: bool = False
    ocr_used: bool = False
    extractor_used: str = "pymupdf"
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        """Join all page texts preserving page breaks."""
        page_texts = [p.text.strip() for p in self.pages if p.text.strip()]
        return "\n\n--- Page Break ---\n\n".join(page_texts)

    @property
    def page_count(self) -> int:
        return self.total_pages

    def to_dict(self) -> dict[str, Any]:
        return {
            "pages": [p.to_dict() for p in self.pages],
            "total_pages": self.total_pages,
            "total_chars": self.total_chars,
            "is_scanned": self.is_scanned,
            "ocr_used": self.ocr_used,
            "extractor_used": self.extractor_used,
            "full_text": self.full_text,
            "metadata": self.metadata,
        }


class DocumentTextExtractor:
    """
    Core document text extraction service.
    Orchestrates native PDF text extraction, OCR fallback, and image extraction.
    """

    def __init__(self, ocr_service: Optional[OCRService] = None):
        self.ocr_service = ocr_service or get_ocr_service()
        self.min_chars_threshold = settings.OCR_MIN_CHARS_THRESHOLD

    def extract_pages(
        self,
        file_bytes: bytes,
        filename: str,
        content_type: str = "application/octet-stream",
    ) -> DocumentPagesResult:
        """
        Extract page-by-page text from supported documents.
        Preserves page boundaries and records whether OCR was used per page.
        """
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if content_type == "application/pdf" or ext == "pdf":
            return self._extract_pdf_pages(file_bytes, filename)
        elif (
            content_type.startswith("image/")
            or ext in {"jpg", "jpeg", "png", "webp", "bmp", "tiff"}
        ):
            return self._extract_image_pages(file_bytes, filename, content_type)
        elif ext == "docx" or "wordprocessingml" in content_type:
            return self._extract_docx_pages(file_bytes, filename)
        elif ext == "xlsx" or "spreadsheetml" in content_type:
            return self._extract_xlsx_pages(file_bytes, filename)
        else:
            return self._extract_plain_text_pages(file_bytes, filename)

    def extract_text(
        self,
        file_bytes: bytes,
        filename: str,
        content_type: str = "application/octet-stream",
    ) -> str:
        """Extract full document text joined across all pages."""
        result = self.extract_pages(file_bytes, filename, content_type)
        return result.full_text

    def _extract_pdf_pages(self, file_bytes: bytes, filename: str) -> DocumentPagesResult:
        """
        Extract pages from PDF.
        Attempts native PyMuPDF text first; falls back to OCR if page has little/no text.
        """
        pages: List[PageData] = []
        any_ocr_used = False
        total_chars = 0
        doc_meta = {}

        try:
            import pymupdf
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            total_pages = len(doc)
            doc_meta = doc.metadata or {}

            for idx in range(total_pages):
                page_num = idx + 1
                page = doc[idx]
                native_text = (page.get_text("text") or "").strip()

                # Clean lines
                lines = [l.strip() for l in native_text.splitlines() if l.strip()]
                clean_text = "\n".join(lines)

                # Check if native text is sufficient or if OCR fallback is needed
                if len(clean_text) >= self.min_chars_threshold:
                    # Native selectable text is good!
                    pages.append(
                        PageData(
                            page_number=page_num,
                            text=clean_text,
                            ocr_used=False,
                            char_count=len(clean_text),
                            lines=lines,
                        )
                    )
                    total_chars += len(clean_text)
                else:
                    # Scanned or image-only page: render to image & run OCR
                    logger.info(
                        "PDF page has insufficient native text; triggering OCR fallback",
                        filename=filename,
                        page_number=page_num,
                        native_chars=len(clean_text),
                    )
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        ocr_text, provider = self.ocr_service.perform_ocr_on_image(img_bytes)
                        ocr_lines = [l.strip() for l in ocr_text.splitlines() if l.strip()]
                        final_text = "\n".join(ocr_lines) if ocr_lines else clean_text

                        pages.append(
                            PageData(
                                page_number=page_num,
                                text=final_text,
                                ocr_used=True,
                                char_count=len(final_text),
                                lines=ocr_lines or lines,
                            )
                        )
                        any_ocr_used = True
                        total_chars += len(final_text)
                    except Exception as ocr_err:
                        logger.warning(
                            "OCR render failed for PDF page, keeping native text",
                            page=page_num,
                            error=str(ocr_err),
                        )
                        pages.append(
                            PageData(
                                page_number=page_num,
                                text=clean_text,
                                ocr_used=False,
                                char_count=len(clean_text),
                                lines=lines,
                            )
                        )
                        total_chars += len(clean_text)

            doc.close()
            is_scanned = any_ocr_used or (total_pages > 0 and (total_chars / total_pages) < 50)

            return DocumentPagesResult(
                pages=pages,
                total_pages=total_pages,
                total_chars=total_chars,
                is_scanned=is_scanned,
                ocr_used=any_ocr_used,
                extractor_used="pymupdf+ocr" if any_ocr_used else "pymupdf",
                metadata={"title": doc_meta.get("title"), "author": doc_meta.get("author")},
            )

        except Exception as e:
            logger.warning("PyMuPDF failed, attempting pypdf fallback", error=str(e), filename=filename)

        # Fallback to pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            total_pages = len(reader.pages)
            for idx, p in enumerate(reader.pages):
                txt = (p.extract_text() or "").strip()
                lines = [l.strip() for l in txt.splitlines() if l.strip()]
                pages.append(
                    PageData(
                        page_number=idx + 1,
                        text="\n".join(lines),
                        ocr_used=False,
                        char_count=len(txt),
                        lines=lines,
                    )
                )
                total_chars += len(txt)

            return DocumentPagesResult(
                pages=pages,
                total_pages=total_pages,
                total_chars=total_chars,
                is_scanned=False,
                ocr_used=False,
                extractor_used="pypdf",
            )
        except Exception as e2:
            logger.error("All PDF extractors failed", error=str(e2), filename=filename)
            return DocumentPagesResult(
                pages=[],
                total_pages=0,
                total_chars=0,
                is_scanned=True,
                ocr_used=False,
                extractor_used="failed",
            )

    def _extract_image_pages(
        self, file_bytes: bytes, filename: str, content_type: str
    ) -> DocumentPagesResult:
        """Extract text from image using OCRService."""
        ocr_text, provider = self.ocr_service.perform_ocr_on_image(file_bytes)
        lines = [l.strip() for l in ocr_text.splitlines() if l.strip()]
        clean_text = "\n".join(lines)

        page = PageData(
            page_number=1,
            text=clean_text,
            ocr_used=True,
            char_count=len(clean_text),
            lines=lines,
        )

        return DocumentPagesResult(
            pages=[page],
            total_pages=1,
            total_chars=len(clean_text),
            is_scanned=True,
            ocr_used=True,
            extractor_used=f"image_ocr_{provider}",
            metadata={"filename": filename, "content_type": content_type},
        )

    def _extract_docx_pages(self, file_bytes: bytes, filename: str) -> DocumentPagesResult:
        """Extract paragraphs and tables from Word documents."""
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            text = "\n".join(paragraphs)
            return DocumentPagesResult(
                pages=[
                    PageData(
                        page_number=1,
                        text=text,
                        ocr_used=False,
                        char_count=len(text),
                        lines=paragraphs,
                    )
                ],
                total_pages=1,
                total_chars=len(text),
                is_scanned=False,
                ocr_used=False,
                extractor_used="python-docx",
            )
        except Exception as e:
            logger.warning("DOCX extraction error", error=str(e), filename=filename)
            return self._extract_plain_text_pages(file_bytes, filename)

    def _extract_xlsx_pages(self, file_bytes: bytes, filename: str) -> DocumentPagesResult:
        """Extract sheet rows from Excel spreadsheets."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            pages: List[PageData] = []
            total_chars = 0

            for idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        rows.append(" | ".join(row_vals))

                sheet_text = "\n".join(rows)
                pages.append(
                    PageData(
                        page_number=idx + 1,
                        text=sheet_text,
                        ocr_used=False,
                        char_count=len(sheet_text),
                        lines=rows,
                    )
                )
                total_chars += len(sheet_text)

            wb.close()
            return DocumentPagesResult(
                pages=pages,
                total_pages=len(pages),
                total_chars=total_chars,
                is_scanned=False,
                ocr_used=False,
                extractor_used="openpyxl",
            )
        except Exception as e:
            logger.warning("XLSX extraction error", error=str(e), filename=filename)
            return self._extract_plain_text_pages(file_bytes, filename)

    def _extract_plain_text_pages(self, file_bytes: bytes, filename: str) -> DocumentPagesResult:
        """Fallback for plaintext or raw bytes."""
        try:
            text = file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            text = ""

        lines = [l.strip() for l in text.splitlines() if l.strip()]
        return DocumentPagesResult(
            pages=[
                PageData(
                    page_number=1,
                    text=text,
                    ocr_used=False,
                    char_count=len(text),
                    lines=lines,
                )
            ],
            total_pages=1,
            total_chars=len(text),
            is_scanned=False,
            ocr_used=False,
            extractor_used="plain_text",
        )
